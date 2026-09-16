import openpyxl,os,time
import uuid
from rich.console import Console
from rich.progress import track
from pyfiglet import figlet_format
from pwinput import pwinput
from datetime import datetime
console=Console()
file_name="bank_records.xlsx"
bank_name=figlet_format("PATREON BANK")





#==================================================================
#*********************PROGRESS BAR*********************************
#==================================================================
def progress_bar(val):
    for _ in track(range(val),description="Processing..."):
        time.sleep(1)



#==========================================================
# Current date and Time 
#==========================================================
def cur_date():
    now = datetime.now()
    # Format: YYYY-MM-DD HH:MM:SS
    formatted_1 = now.strftime("%Y-%m-%d %H:%M:%S")
    return formatted_1
    #print("Format 1:", formatted_1)  # Output: 2026-09-05 22:00:00

#==================================================================
#=================ERROR MSG========================================
#==================================================================
def err_msg(error):
    console.print(f"[bold red]INVALID {error} TRY AGAIN[bold red]")


#==================================================================
#*****************EXCEL CREATE FUNCTION****************************
#==================================================================

def create_excel_file():
    """This function creates a excel file named bank_records.xlsx If it doesn't exist"""
    if not os.path.exists(file_name):
        wb=openpyxl.Workbook()
        sheet=wb.active
        sheet.title="Bank Records"
        header=["Account No.","Name","PIN","Transaction ID","Transaction Type","Amount","Previous Balance","Current Balance","Date-Time"]
        sheet.append(header)

        wb.save(file_name)
        wb.close()
        console.print("[bold blue]Excel Database Created Successfully[/bold blue]")

#==================================================================
#************************CREATE ACCOUNT****************************
#==================================================================
def create_acc():
    console.print("[bold yellow]Create Account[/bold yellow]")
    name = input("Enter your name: ").strip().title()
    acc_no = input("Enter Account Number: ").strip()
    if not acc_no.isdigit():
        err_msg("ACCOUNT NUMBER")
        return
    wb=openpyxl.load_workbook(file_name)
    sheet=wb.active
    for row in sheet.iter_rows(min_row=2,values_only=True):
        if acc_no==row[0]:
            err_msg("(ACCOUNT NUMBER ALREADY EXISTS)")
            return
        if len(acc_no)!=10:
            err_msg("(ACCOUNT NUMBER MUST BE 10 DIGITS)")
            return
    pin=pwinput("CREATE A 6 DIGIT PIN: ",mask="*")

    if len(pin)!=6:
        err_msg("PIN")
        return
    try:
        amount=float(input("Enter opening balance: "))
        if amount<0:
            err_msg("OPENING BALANCE")
            return
        transaction_id=str(uuid.uuid4())[:10]
        status="Opening"
        dt=cur_date()
        sheet.append([acc_no,name,pin,transaction_id,status,amount,0,amount,dt])
    except ValueError:
        err_msg("OPENING BALANCE")
    wb.save(file_name)
    wb.close()
    progress_bar(5)

    console.print("[bold green]Account Created Successfully[/bold green]")

#==================================================================
#==============================LOGIN===============================
#==================================================================
def acc_login():
    console.print("[bold blackon white]LOGIN[bold black on white]")

    account_no=input("Enter account number: ").strip()
    
    pin=pwinput("Enter PIN: ", mask="*")
    wb = openpyxl.load_workbook(file_name)
    sheet = wb.active
    for row in sheet.iter_rows(min_row=2,min_col=1,values_only=True):
        user_name=row[1].title()
        
        
        if str(row[0]) == str((account_no)) and str(row[2]) == str(pin):
            console.print("[yellow on black]Verifying account.....[/yellow on black]")
            progress_bar(1)
            console.print("[bold green]Login Successful[/bold green]")
            console.print(f"[white]Welcome {user_name}[/white]")
            sub_menu(account_no)
            break       
    else:
        console.print("[yellow on black]Verifying details.....[/yellow on black]")
        progress_bar(5)
        err_msg("Account id or PIN")
    wb.close()
def get_user_details(account_no):
    wb=openpyxl.load_workbook(file_name)
    sheet=wb.active
    name=None
    pin=None
    for row in sheet.iter_rows(min_row=2,values_only=True):
        if str(row[0])==str(account_no):
            name=row[1]
            pin=row[0]
            break
    wb.close()   
    return name,pin

def get_balance(account_no):
    wb=openpyxl.load_workbook(file_name)
    sheet=wb.active
    current_balance=None
    for row in sheet.iter_rows(min_row=2,values_only=True):
        if str(row[0])==str(account_no):
            current_balance=row[-2]
    wb.close()
    return current_balance
            


#==================================================================
#************************CHECK BALANCE*****************************
#==================================================================
def check_bal(account_no):
    balance=get_balance(account_no)
    console.print(f"\n[bold green]Current Balance: ${balance:.2f}[bold green]")



#==================================================================
#============================DEPOSIT===============================
#==================================================================
def deposit(account_no):
    console.print(f"[bold magenta]Deposit Money[bold magenta]".center(100))

    
    try:
        amount= float(input("Enter the amount of money to deposit: "))
        if amount<0 or amount==0:
            err_msg("DEPOSIT BALANCE")
            return
    except ValueError:
        err_msg("DEPOSIT BALANCE")
        return
    
    dt=cur_date()

    transaction_id=str(uuid.uuid4())[:10]

    name,pin=get_user_details(account_no)

    if name==None:
        console.print("[red]Corrupted data name doesn't exist but account does.[red]")
        return
    
    transaction_type="Deposit"

    prev_bal=get_balance(account_no)

    current_bal=prev_bal+amount

    wb=openpyxl.load_workbook(file_name)
    sheet=wb.active
    sheet.append([account_no,name,None,transaction_id,transaction_type,amount,prev_bal,current_bal,dt])
    
    progress_bar(10)
    console.print(f"[bold green on white]Money deposited successfully \nCurrent Balance: ${current_bal}[/bold green on white]".center(100))
    wb.save(file_name)
    wb.close()
    
        
#==================================================================
#===========================SUB MENU===============================
#==================================================================
def sub_menu(account_no):
    while True:
        wb=openpyxl.load_workbook(file_name)
        sheet=wb.active
        console.print("""[white on black]   Banking Options      
1. Check Balance        
2. Deposit Money        
3. Withdraw Money       
4. Transaction History  
5. Logout               [/white on black]
    """)
        choice = input("Enter Choice: ").strip()
        if choice=="1":
            check_bal(account_no)
        elif choice=="2":
            deposit(account_no)
        elif choice=="3":
            pass
        elif choice=="4":
            pass
        elif choice=="5":
            break
        else:
            err_msg("CHOICE")




#==================================================================
#===========================MAIN PROGRAM===========================
#==================================================================


create_excel_file()
while True:
    console.print(f"[green]{bank_name}[/green]")
    print("""1. Create Account
2. Login
3. Exit""")
    print()
    choice=console.input("[bold magenta]Enter choice: [/bold magenta]").strip()
    if choice=="1":
        create_acc()
    elif choice=="2":
        acc_login()
    elif choice=="3":
        break
    else:
        console.print("[red]Invalid choice entered[/red]")

        

